"""Import PowerPoint (.pptx) decks into Lesson/Slide records.

Each slide becomes a Slide row: the first text frame is treated as the
title, remaining text frames become the body, and speaker notes are kept
(useful later for audio narration). Images are extracted to ``static/slides``.

Usage (CLI):
    python -m app.pptx_importer path/to/deck.pptx --title "Saving Basics"
"""

from __future__ import annotations

import argparse
import os
from pathlib import Path

from pptx import Presentation
from sqlalchemy.orm import Session

from app.database import SessionLocal, init_db
from app.models import AgeBand, Lesson, Slide
from app.translation import Translator

STATIC_DIR = Path("static/slides")


def _extract_slide_text(slide) -> tuple[str, str]:
    """Return (title, body) text for a slide."""
    title = ""
    body_parts: list[str] = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text.strip()
        if not text:
            continue
        if not title:
            title = text.splitlines()[0]
            rest = "\n".join(text.splitlines()[1:]).strip()
            if rest:
                body_parts.append(rest)
        else:
            body_parts.append(text)
    return title, "\n\n".join(body_parts)


def _extract_first_image(slide, lesson_slug: str, index: int) -> str | None:
    """Save the first picture on a slide and return its relative path."""
    STATIC_DIR.mkdir(parents=True, exist_ok=True)
    for shape in slide.shapes:
        if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
            image = shape.image
            ext = image.ext or "png"
            filename = f"{lesson_slug}_{index}.{ext}"
            path = STATIC_DIR / filename
            with open(path, "wb") as fh:
                fh.write(image.blob)
            return str(path).replace("\\", "/")
    return None


def import_pptx(
    db: Session,
    file_path: str,
    title: str | None = None,
    age_band: AgeBand = AgeBand.TEENS,
    locale: str = "de-CH",
    translate_from: str | None = None,
) -> Lesson:
    prs = Presentation(file_path)

    # Optionally translate slide text from the source language (e.g. Spanish)
    # into the target launch locale before persisting.
    translator = Translator(
        target_locale=locale,
        source=translate_from or "auto",
        enabled=translate_from is not None,
    )

    # Only translate a title we derived from the filename; an explicitly
    # provided title is assumed to already be in the target language.
    lesson_title = title or Path(file_path).stem
    if translate_from is not None and title is None:
        lesson_title = translator.translate(lesson_title)
    lesson_slug = "".join(c for c in lesson_title.lower() if c.isalnum() or c == "_")[:40]

    lesson = Lesson(
        title=lesson_title,
        age_band=age_band,
        locale=locale,
        source_file=os.path.basename(file_path),
    )
    db.add(lesson)
    db.flush()

    # First pass: extract all slide text + images.
    raw_slides: list[dict] = []
    for i, pptx_slide in enumerate(prs.slides):
        s_title, s_body = _extract_slide_text(pptx_slide)
        image_path = _extract_first_image(pptx_slide, lesson_slug, i)
        notes = ""
        if pptx_slide.has_notes_slide:
            notes = pptx_slide.notes_slide.notes_text_frame.text.strip()
        raw_slides.append(
            {"title": s_title, "body": s_body, "image_path": image_path, "notes": notes}
        )

    # Second pass: translate every text field in one batched operation
    # (packs many segments per request → avoids rate-limit bans).
    if translate_from is not None:
        flat = [s[field] for s in raw_slides for field in ("title", "body", "notes")]
        translated = translator.translate_batch(flat)
        for idx, s in enumerate(raw_slides):
            s["title"], s["body"], s["notes"] = translated[idx * 3 : idx * 3 + 3]

    for i, s in enumerate(raw_slides):
        db.add(
            Slide(
                lesson_id=lesson.id,
                order_index=i,
                title=s["title"],
                body=s["body"],
                image_path=s["image_path"],
                notes=s["notes"],
            )
        )

    db.commit()
    db.refresh(lesson)
    return lesson


def _main() -> None:
    parser = argparse.ArgumentParser(description="Import a PPTX deck as a lesson.")
    parser.add_argument("file", help="Path to the .pptx file")
    parser.add_argument("--title", default=None, help="Lesson title")
    parser.add_argument(
        "--age-band",
        default=AgeBand.TEENS.value,
        choices=[b.value for b in AgeBand],
    )
    parser.add_argument("--locale", default="de-CH")
    parser.add_argument(
        "--translate-from",
        default=None,
        help=(
            "Source language code of the deck (e.g. 'es' for Spanish). When set, "
            "slide titles, bodies and notes are translated into --locale before import."
        ),
    )
    args = parser.parse_args()

    init_db()
    db = SessionLocal()
    try:
        lesson = import_pptx(
            db,
            args.file,
            title=args.title,
            age_band=AgeBand(args.age_band),
            locale=args.locale,
            translate_from=args.translate_from,
        )
        print(f"Imported lesson '{lesson.title}' with {len(lesson.slides)} slides (id={lesson.id}).")
    finally:
        db.close()


if __name__ == "__main__":
    _main()

